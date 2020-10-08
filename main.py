from docx import Document

doc = Document('SD-App 20201008.docx')

numTables = doc.tables
table = numTables[0]

table.cell(3, 3).text = "Michelle"

doc.save('SD-App 20201008.docx')
# table.cell(3, 4).text = "Michelle"

# for table in numTables:
#     row_count = len(table.rows)
#     col_count = len(table.columns)
#     for i in range(row_count):
#         for j in range(col_count):
#             if "2020/10/08" in table.cell(i, j).text:
#                 print(i, j)
#             else:
#                 continue