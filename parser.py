from docx import Document
import json
import os
import Constant


def file_renaming(fname):

    if 'SD-App' not in fname:
        origin_fname = fname.split('.')[0]
        year = origin_fname.split('第')[0]
        date = ''
        if '~' in origin_fname:
            date = origin_fname.split('~')[1].replace('/', '').replace(')', '').strip()
        else:
            dates = origin_fname.split('_')
            date = dates[2].strip() + dates[3].replace(')', '').strip()

        new_fname = 'SD-App ' + year + date
        return new_fname
        #os.rename(plain_text_path+fname, plain_text_path+new_fname)

def addWord(s, data):
    if 'marc' in data:
        data[Constant.marc_key] += s
    else:
        data[Constant.hunk_key] += s
    data[Constant.william_key] += s
    data[Constant.michelle_key] += s
    data[Constant.sandy_key] += s
    data[Constant.eric_key] += s


def parse_person(parts, data):
    # if part == weeklyComplete:
    #     addWord('本週完成項目', data)
    # elif part == processing:
    #     addWord('進行中項目', data)
    # else:
    #     addWord('下週預計', data)
    for part in parts:
        if parts.index(part) == 0:
            addWord('本週完成項目', data)
        elif parts.index(part) == 1:
            addWord('進行中項目', data)
        else:
            addWord('下週預計', data)
        
        if 'marc' in data:
            data[Constant.marc_key] += part.split(Constant.marc)[1].split(Constant.william)[0]
        else:
            data[Constant.hunk_key] += part.split(Constant.hunk)[1].split(Constant.william)[0]

        data[Constant.william_key] += part.split(Constant.william)[1].split(Constant.michelle)[0]
        data[Constant.michelle_key] += part.split(Constant.michelle)[1].split('Android')[0]
        data[Constant.sandy_key] += part.split(Constant.sandy)[1].split(Constant.eric)[0]
        data[Constant.eric_key] += part.split(Constant.eric)[1]


def parse_plain_text(fname):
    data = {Constant.william_key:'', Constant.michelle_key:'', Constant.sandy_key:'', Constant.eric_key:''}
    
    plain_text_path = 'plainText/'

    f = open(plain_text_path + '/' +fname, 'r')

    content = f.read()

    if Constant.marc in content:
        data[Constant.marc_key] = ''
    elif Constant.hunk in content:
        data[Constant.hunk_key] = ''

    weeklyComplete = content.split('本週完成項目：')[1].split('進行中項目：')[0]
    processing = content.split('進行中項目：')[1].split('下週預計：')[0]
    futurePlan = content.split('下週預計：')[1]

    contents = [weeklyComplete, processing, futurePlan]

    # parse_person(weeklyComplete)
    # parse_person(processing)
    # parse_person(futurePlan)
    parse_person(contents, data)


    f.close()
    #parse 中文
    json_data = json.dumps(data,ensure_ascii=False,indent=2)

    new_fname = file_renaming(fname)

    new_path = 'json/'

    with open(new_path + new_fname +'.json', "w") as outfile: 
        outfile.write(json_data) 


def write_to_word(fname):

#sample Documents
    doc = Document('SD-App 20201008.docx')

    numTables = doc.tables
    table = numTables[0]

#會議日期
    date = fname.split('SD-App')[1].split('.')[0].strip()
    table.cell(1, 1).text = date

# 紀錄
    # table.cell(3, 3).text = "Michelle"

#read file
    first_filename = fname.split('.')[0]
    json_folder_path = 'json/'

    with open(json_folder_path + fname , 'r') as reader:
        json_data = json.loads(reader.read())
    #主席
        if Constant.marc_key in json_data:
            table.cell(2, 1).text = 'Marc'
        else:
            table.cell(2, 1).text = 'Hunk'

##########Leader
#工作項目
        if Constant.marc_key in json_data:
            table.cell(7, 4).text = 'Marc'
            table.cell(7, 0).text = json_data[Constant.marc_key]
        else:
            table.cell(7, 4).text = 'Hunk'
            table.cell(7, 0).text = json_data[Constant.hunk_key]

######william
#工作項目
        table.cell(8, 0).text = json_data[Constant.william_key]

######michelle
#工作項目
        table.cell(9, 0).text = json_data[Constant.michelle_key]

######sandy
#工作項目
        table.cell(10, 0).text = json_data[Constant.sandy_key]

######eric
#工作項目
        table.cell(11, 0).text = json_data[Constant.eric_key]

    results_path = 'results/'
    doc.save(results_path+first_filename+'.docx')


    



#file_renaming('2020第01週 App Team 週報 (12_30 _ 01_03).docx')